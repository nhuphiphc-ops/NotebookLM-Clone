"""
Quét kho tài liệu (docs/) và chuẩn hoá về định dạng mà Gemini File API đọc được.

- .pdf / .txt / .md / .csv  -> nạp trực tiếp
- .docx / .xlsx             -> chuyển sang Markdown trong processed_docs/
- Các định dạng khác        -> bỏ qua và báo lý do

Hỗ trợ thư mục con: tên tài liệu là đường dẫn tương đối so với docs/
(ví dụ: "Hop dong/HD-2026.pdf") nên trích dẫn của AI vẫn giữ được thư mục.
"""

import mimetypes
import os
import re

DOCS_DIR = 'docs'
PROCESSED_DOCS_DIR = 'processed_docs'

# Nạp thẳng lên Gemini, không cần chuyển đổi
DIRECT_EXTS = {'.pdf', '.txt', '.md', '.csv'}
# Cần chuyển sang Markdown trước khi nạp
CONVERT_EXTS = {'.docx', '.xlsx'}
SUPPORTED_EXTS = DIRECT_EXTS | CONVERT_EXTS

# KHONG dung mimetypes.guess_type() lam nguon chinh: tren Windows, registry
# thuong map .csv -> application/vnd.ms-excel, va Gemini tu choi MIME do
# (400 INVALID_ARGUMENT: Unsupported MIME type).
MIME_TYPES = {
    '.pdf': 'application/pdf',
    '.txt': 'text/plain',
    '.md': 'text/markdown',
    '.csv': 'text/csv',
}

# Giới hạn CỨNG của Gemini với tài liệu PDF: 50 MB hoặc 1000 trang.
# File vuot nguong van nap len File API duoc nhung se loi o buoc hoi dap,
# nen phai chan tu vong quet de bao som cho nguoi dung.
PDF_MAX_BYTES = 50 * 1024 * 1024
PDF_MAX_PAGES = 1000

# Định dạng cũ / không đọc được -> thông báo cụ thể cho người dùng
UNSUPPORTED_HINTS = {
    '.doc': 'Định dạng .doc cũ — hãy lưu lại thành .docx',
    '.xls': 'Định dạng .xls cũ — hãy lưu lại thành .xlsx',
    '.ppt': 'Định dạng PowerPoint chưa được hỗ trợ',
    '.pptx': 'Định dạng PowerPoint chưa được hỗ trợ',
    '.zip': 'Hãy giải nén trước khi nạp',
    '.rar': 'Hãy giải nén trước khi nạp',
}


def get_mime_type(file_path):
    """MIME type an toàn cho Gemini File API."""
    ext = os.path.splitext(file_path)[1].lower()
    if ext in MIME_TYPES:
        return MIME_TYPES[ext]
    guessed = mimetypes.guess_type(file_path)[0]
    return guessed or 'text/plain'


def _is_ignorable(filename):
    """File rác / file tạm của Office cần bỏ qua khi quét."""
    if filename.startswith(('~$', '.', '$')):
        return True
    return filename.lower().endswith(('.tmp', '.lnk', '.crdownload', '.part'))


def _ensure_dirs():
    os.makedirs(DOCS_DIR, exist_ok=True)
    os.makedirs(PROCESSED_DOCS_DIR, exist_ok=True)


def _rel_name(abs_path):
    """Tên tài liệu = đường dẫn tương đối so với docs/, luôn dùng dấu '/'."""
    rel = os.path.relpath(abs_path, DOCS_DIR)
    return rel.replace(os.sep, '/')


def _processed_path(rel_name):
    """Đường dẫn file Markdown trung gian, phẳng hoá và làm sạch tên."""
    safe = re.sub(r'[^0-9A-Za-z._-]+', '_', rel_name.replace('/', '__'))
    base = os.path.splitext(safe)[0][:120]
    return os.path.join(PROCESSED_DOCS_DIR, base + '.md')


def processed_path(rel_name):
    """Công khai quy tắc đặt tên bản chuyển đổi để app.py xoá được đúng file."""
    return _processed_path(rel_name)


def _pdf_over_limit(abs_path, size):
    """Trả về lý do nếu PDF vượt giới hạn của Gemini, ngược lại trả về ''."""
    if size > PDF_MAX_BYTES:
        return (
            'PDF nặng %.0f MB, vượt giới hạn 50 MB của Gemini — hãy tách nhỏ file'
            % (size / 1024 / 1024)
        )
    try:
        import pypdf

        pages = len(pypdf.PdfReader(abs_path, strict=False).pages)
    except Exception:
        return ''  # đọc không được thì cứ để Gemini phán, không chặn oan
    if pages > PDF_MAX_PAGES:
        return 'PDF có %d trang, vượt giới hạn 1000 trang của Gemini — hãy tách nhỏ file' % pages
    return ''


def list_documents():
    """
    Liệt kê mọi tài liệu trong docs/ (bao gồm thư mục con).

    Trả về list dict: name, path, folder, ext, size, mtime, supported, reason.
    """
    _ensure_dirs()
    docs = []
    for root, dirnames, filenames in os.walk(DOCS_DIR):
        dirnames[:] = sorted(d for d in dirnames if not _is_ignorable(d))
        for filename in sorted(filenames):
            if _is_ignorable(filename):
                continue
            abs_path = os.path.join(root, filename)
            if not os.path.isfile(abs_path):
                continue
            rel = _rel_name(abs_path)
            ext = os.path.splitext(filename)[1].lower()
            supported = ext in SUPPORTED_EXTS
            reason = ''
            size = os.path.getsize(abs_path)
            if not supported:
                reason = UNSUPPORTED_HINTS.get(
                    ext, 'Định dạng ' + (ext or 'không rõ') + ' chưa được hỗ trợ'
                )
            elif size == 0:
                supported, reason = False, 'File rỗng (0 byte)'
            elif ext == '.pdf':
                over = _pdf_over_limit(abs_path, size)
                if over:
                    supported, reason = False, over
            docs.append({
                'name': rel,
                'path': abs_path,
                'folder': os.path.dirname(rel),
                'ext': ext,
                'size': size,
                'mtime': os.path.getmtime(abs_path),
                'supported': supported,
                'reason': reason,
            })
    return docs


def list_folders():
    """Thư mục con trực tiếp của docs/ kèm số tài liệu ('' = thư mục gốc)."""
    counts = {}
    for doc in list_documents():
        top = doc['folder'].split('/')[0] if doc['folder'] else ''
        counts[top] = counts.get(top, 0) + 1
    _ensure_dirs()
    for entry in sorted(os.listdir(DOCS_DIR)):
        if os.path.isdir(os.path.join(DOCS_DIR, entry)) and not _is_ignorable(entry):
            counts.setdefault(entry, 0)
    return counts


def _title_block(title):
    """Ghi tên tài liệu gốc vào đầu bản chuyển đổi để AI trích dẫn đúng nguồn."""
    if not title:
        return []
    return ['> Tài liệu nguồn: **' + title + '**']


def process_docx(file_path, output_path, title=None):
    """DOCX -> Markdown, giữ cấp tiêu đề và bảng."""
    from docx import Document

    doc = Document(file_path)
    lines = _title_block(title)
    for para in doc.paragraphs:
        text = para.text.strip()
        if not text:
            continue
        style = (para.style.name or '') if para.style is not None else ''
        match = re.match(r'Heading (\d)', style)
        if match:
            lines.append('#' * min(int(match.group(1)), 6) + ' ' + text)
        else:
            lines.append(text)

    for idx, table in enumerate(doc.tables, start=1):
        lines.append('### Bảng ' + str(idx))
        # các dòng của bảng phải liền nhau để Markdown hiểu là một bảng
        rows = []
        for r_idx, row in enumerate(table.rows):
            cells = [c.text.replace('\n', ' ').strip() for c in row.cells]
            rows.append('| ' + ' | '.join(cells) + ' |')
            if r_idx == 0:
                rows.append('| ' + ' | '.join(['---'] * len(cells)) + ' |')
        if rows:
            lines.append('\n'.join(rows))

    _write_text(output_path, '\n\n'.join(lines))
    return output_path


def process_xlsx(file_path, output_path, title=None):
    """XLSX -> Markdown, mỗi sheet một mục để AI trích dẫn được tên sheet."""
    import pandas as pd

    parts = _title_block(title)
    with pd.ExcelFile(file_path) as xl:
        for sheet_name in xl.sheet_names:
            df = xl.parse(sheet_name)
            parts.append('## Sheet: ' + str(sheet_name))
            if df.empty:
                parts.append('_(Sheet rỗng)_')
                continue
            df = df.dropna(axis=1, how='all').dropna(axis=0, how='all')
            try:
                parts.append(df.to_markdown(index=False))
            except ImportError:
                # thieu 'tabulate' -> van xuat duoc noi dung
                parts.append('```\n' + df.to_string(index=False) + '\n```')

    _write_text(output_path, '\n\n'.join(parts))
    return output_path


def _write_text(output_path, text):
    os.makedirs(os.path.dirname(output_path) or '.', exist_ok=True)
    with open(output_path, 'w', encoding='utf-8') as f:
        f.write(text)


def scan_and_process_documents(uploaded_files_registry):
    """
    So sánh docs/ với registry đã nạp và trả về (files_to_upload, skipped).

    files_to_upload: [{original_name, path, mime_type, mtime, size, folder}]
      - 'path' là file sẽ nạp lên (bản gốc hoặc bản .md đã chuyển đổi)
      - 'original_name' là tên hiển thị/trích dẫn (đường dẫn tương đối)
    skipped: [{name, reason}] để hiển thị cho người dùng, không làm treo tiến trình.
    """
    _ensure_dirs()
    files_to_upload, skipped = [], []

    for doc in list_documents():
        name, mtime = doc['name'], doc['mtime']

        if not doc['supported']:
            skipped.append({'name': name, 'reason': doc['reason']})
            continue

        known = uploaded_files_registry.get(name)
        if known and known.get('mtime') == mtime:
            continue  # đã nạp và chưa thay đổi

        try:
            if doc['ext'] in DIRECT_EXTS:
                upload_path = doc['path']
            else:
                upload_path = _processed_path(name)
                stale = (
                    not os.path.exists(upload_path)
                    or os.path.getmtime(upload_path) < mtime
                )
                if stale:
                    if doc['ext'] == '.docx':
                        process_docx(doc['path'], upload_path, title=name)
                    else:
                        process_xlsx(doc['path'], upload_path, title=name)
        except Exception as exc:  # file hỏng / có mật khẩu / đang bị mở
            skipped.append({'name': name, 'reason': 'Không đọc được: ' + str(exc)})
            continue

        files_to_upload.append({
            'original_name': name,
            'path': upload_path,
            'mime_type': get_mime_type(upload_path),
            'mtime': mtime,
            'size': doc['size'],
            'folder': doc['folder'],
        })

    return files_to_upload, skipped
